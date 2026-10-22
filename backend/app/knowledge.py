"""Document parsing, immutable version chunks, and transactional publishing."""
import io
import re
from datetime import date

from sqlalchemy import select, update

from .db import Chunk, Document, Meta, uid


def parse_file(filename: str, content: bytes) -> str:
    suffix = filename.rsplit('.', 1)[-1].lower()
    if suffix in ('md', 'txt'):
        try:
            return content.decode('utf-8-sig')
        except UnicodeDecodeError:
            raise ValueError('文本文件请使用 UTF-8 编码')
    if suffix == 'docx':
        from docx import Document as WordDocument
        doc = WordDocument(io.BytesIO(content))
        parts = []
        for p in doc.paragraphs:
            if not p.text.strip():
                continue
            style = p.style.name if p.style else ''
            heading = re.search(r'(?:Heading|标题)\s*(\d+)', style, re.I)
            parts.append(('#' * min(int(heading[1]), 6) + ' ' if heading else '') + p.text)
        for table in doc.tables:
            parts.extend(' | '.join(c.text for c in row.cells) for row in table.rows)
        return '\n\n'.join(parts)
    if suffix == 'pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise ValueError('暂不支持加密 PDF')
        pages = [(i + 1, (p.extract_text() or '').strip()) for i, p in enumerate(reader.pages)]
        if not pages or any(not text for _, text in pages):
            raise ValueError('检测到无文字页面；暂不支持扫描件或混合扫描 PDF')
        return '\n\n'.join(f'## PDF 第 {i} 页\n\n{text}' for i, text in pages)
    raise ValueError('仅支持 Markdown、TXT、DOCX 和文字型 PDF')


def add_chunks(db, doc, max_chars=2200):
    if max_chars < 1:
        raise ValueError('分块长度必须为正数')
    section, buffer, n = '正文', [], 0

    def flush():
        nonlocal n
        text = '\n'.join(buffer).strip()
        if text:
            n += 1
            db.add(Chunk(id=uid(), document_id=doc.id, section=section[:300], text=text,
                         location=f'段落组 {n}', refs=sorted(set(re.findall(r'\b[A-Z]+-\d{3}\b', text)))))
        buffer.clear()

    for line in doc.content.splitlines():
        if re.match(r'^#{1,6}\s+', line):
            flush()
            section = re.sub(r'^#{1,6}\s+', '', line).strip()
        else:
            if len(line) > 12000:
                raise ValueError('单个段落过长，请按完整条款添加段落或标题后重试')
            if len('\n'.join(buffer)) + len(line) > max_chars and buffer:
                flush()
            buffer.append(line)
    flush()
    db.flush()
    return n


def validate_dates(start, end):
    date.fromisoformat(start)
    if end:
        date.fromisoformat(end)
        if end <= start:
            raise ValueError('失效日期必须晚于生效日期，区间为 [生效日期, 失效日期)')


def publish(db, doc, indexer, replace_previous=False):
    # Serialize publications before checking overlap and exposing either index.
    db.execute(update(Meta).where(Meta.key == 'revision').values(value=Meta.value + 0))
    validate_dates(doc.effective_from, doc.effective_to)
    if doc.status == 'published':
        return
    others = db.scalars(select(Document).where(Document.code == doc.code, Document.status == 'published', Document.id != doc.id)).all()
    for old in others:
        if doc.effective_from < (old.effective_to or '9999-12-31') and old.effective_from < (doc.effective_to or '9999-12-31'):
            if replace_previous and old.effective_from < doc.effective_from and (doc.effective_to or '9999-12-31') >= (old.effective_to or '9999-12-31'):
                old.effective_to = doc.effective_from
            else:
                raise ValueError(f'与已发布版本 {old.version} 的有效期重叠。可使用替换发布，在新版生效日关闭旧版本有效期')
    chunks = db.scalars(select(Chunk).where(Chunk.document_id == doc.id)).all()
    if not chunks:
        raise ValueError('文档没有可检索内容')
    # Vectors may be staged before publication; all search results are checked
    # against PostgreSQL visibility so a partial upsert cannot expose a draft.
    indexer.index_document(doc, chunks)
    doc.status = 'published'
    db.execute(update(Meta).where(Meta.key == 'revision').values(value=Meta.value + 1))


def serialize_document(d, include_content=False):
    result = {k: getattr(d, k) for k in ('id', 'code', 'title', 'version', 'category', 'effective_from', 'effective_to', 'status', 'created_at')}
    if include_content:
        result['content'] = d.content
    return result
